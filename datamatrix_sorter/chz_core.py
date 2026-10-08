"""
Ядро: чтение файлов (txt/csv/xlsx), разбор кодов «Честного знака»,
очистка от брака и дублей, сравнение массивов М1 и М2, экспорт.
"""
import os
import re
from collections import Counter, OrderedDict
from dataclasses import dataclass

from datamatrix import GS

START = re.compile(r"(?=01\d{14}21)")


# ---------------------------------------------------------------- модель
@dataclass(frozen=True)
class Code:
    gtin: str       # 14 цифр
    serial: str     # серийный номер (AI 21)
    tail: str       # криптохвост: "93xxxx" или "91xxxx92<подпись>"

    @property
    def kind(self):
        return "длинный (91/92)" if self.tail.startswith("91") else "короткий (93)"

    @property
    def key(self):
        """Код без криптохвоста: 01 + GTIN + 21 + серийный номер."""
        return "01%s21%s" % (self.gtin, self.serial)

    @property
    def full(self):
        """Код с криптохвостом, без разделителей."""
        return self.key + self.tail

    @property
    def gs_text(self):
        """Текст для DataMatrix: GS (0x1D) после полей переменной длины."""
        if self.tail.startswith("91"):
            return "%s%s%s%s%s" % (self.key, GS, self.tail[:6], GS, self.tail[6:])
        return "%s%s%s" % (self.key, GS, self.tail)


@dataclass
class Defect:
    source: str
    where: str
    raw: str
    reason: str
    fatal: bool = True   # False = предупреждение (код принят, но с поправкой)


class _Bad(Exception):
    pass


# ---------------------------------------------------------------- разбор
def gtin_ok(g):
    if len(g) != 14 or not g.isdigit():
        return False
    s = sum(int(d) * (3 if i % 2 == 0 else 1) for i, d in enumerate(g[:13]))
    return (10 - s % 10) % 10 == int(g[13])


def _try_piece(p):
    if any(not (33 <= ord(ch) <= 126) for ch in p):
        raise _Bad("недопустимый символ (ошибка считывания)")
    gtin = p[2:16]
    if len(p) >= 72 and p[-46:-44] == "92" and p[-52:-50] == "91":
        serial, tail = p[18:-52], p[-52:]
    elif len(p) >= 25 and p[-6:-4] == "93":
        serial, tail = p[18:-6], p[-6:]
    else:
        raise _Bad("неполный или повреждённый код (нет криптохвоста)")
    if not 1 <= len(serial) <= 20:
        raise _Bad("неверная длина серийного номера")
    if not gtin_ok(gtin):
        raise _Bad("неверная контрольная цифра GTIN")
    return Code(gtin, serial, tail)


def parse_cell(raw):
    """Одна ячейка/строка -> (список Code, список (причина, фатальность, фрагмент))."""
    s = raw.replace(GS, "").replace("\r", "").replace("\n", "").replace("\ufeff", "").strip()
    s = re.sub(r"^\]d2", "", s)
    if not s:
        return [], []
    if s.startswith('"'):                     # CSV-экранирование
        s = s[1:]
        if s.endswith('"'):
            s = s[:-1]
        s = s.replace('""', '"')
    starts = [m.start() for m in START.finditer(s)]
    if not starts:
        return [], [("не код DataMatrix (нет 01+GTIN+21): штрихкод, ссылка или обрывок", True, s)]
    codes, issues = [], []
    if starts[0] > 0:
        issues.append(("лишние символы перед кодом отброшены: «%s»" % s[:starts[0]][:30], False, s))
    bounds = starts + [len(s)]
    for a, b in zip(bounds, bounds[1:]):
        piece = s[a:b]
        try:
            codes.append(_try_piece(piece))
        except _Bad as e1:
            m = re.search(r"\d{13}$", piece)      # приклеенный штрихкод EAN-13 в конце
            if m and len(piece) > 13:
                try:
                    codes.append(_try_piece(piece[:-13]))
                    issues.append(("приклеенный штрихкод в конце отброшен", False, piece))
                    continue
                except _Bad:
                    pass
            issues.append((str(e1), True, piece))
    return codes, issues


# ---------------------------------------------------------------- чтение файлов
def _read_text_lines(path):
    data = open(path, "rb").read()
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("latin-1")
    for i, line in enumerate(text.split("\n"), 1):
        line = line.rstrip("\r")
        if line.strip():
            yield "строка %d" % i, line


def _read_xlsx(path):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        for r, row in enumerate(ws.iter_rows(values_only=True), 1):
            for c, val in enumerate(row, 1):
                if val is None or val == "":
                    continue
                if isinstance(val, float) and val.is_integer():
                    val = int(val)
                yield "%s!%s%d" % (ws.title, get_column_letter(c), r), str(val)
    wb.close()


def read_cells(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx" or ext == ".xlsm":
        return _read_xlsx(path)
    if ext in (".txt", ".csv", ".tsv", ".log", ""):
        return _read_text_lines(path)
    raise ValueError("Неподдерживаемый формат файла: %s" % ext)


# ---------------------------------------------------------------- набор данных
class Dataset:
    def __init__(self, name):
        self.name = name
        self.codes = OrderedDict()      # key -> Code (без дублей)
        self.defects = []
        self.files = []                 # (имя, ячеек, принято, дублей, брака)
        self.cells = 0
        self.found = 0                  # найдено кодов до удаления дублей
        self.duplicates = 0

    @property
    def bad(self):
        return sum(1 for d in self.defects if d.fatal)


def build_dataset(name, paths, progress=None):
    ds = Dataset(name)
    for path in paths:
        base = os.path.basename(path)
        n_cells = ok = dup = bad = 0
        for where, raw in read_cells(path):
            n_cells += 1
            codes, issues = parse_cell(raw)
            for reason, fatal, frag in issues:
                ds.defects.append(Defect(base, where, frag, reason, fatal))
                if fatal:
                    bad += 1
            for c in codes:
                ok += 1
                if c.key in ds.codes:
                    dup += 1
                else:
                    ds.codes[c.key] = c
            if progress and n_cells % 500 == 0:
                progress("%s: %s — %d" % (name, base, n_cells))
        ds.files.append((base, n_cells, ok, dup, bad))
        ds.cells += n_cells
        ds.found += ok
        ds.duplicates += dup
    return ds


def compare(m1, m2):
    """М3 = коды из М1, которых нет в М2 (по коду без криптохвоста). Хвост берётся из М1."""
    m3 = [c for k, c in m1.codes.items() if k not in m2.codes]
    only_m2 = [c for k, c in m2.codes.items() if k not in m1.codes]
    both = len(m1.codes) - len(m3)
    return m3, both, only_m2


def sort_codes(codes, mode="gtin"):
    if mode == "serial":
        return sorted(codes, key=lambda c: (c.serial, c.gtin))
    if mode == "code":
        return sorted(codes, key=lambda c: c.key)
    if mode == "source":
        return list(codes)
    return sorted(codes, key=lambda c: (c.gtin, c.serial))


def filter_codes(codes, text):
    """Отбор по части кода: GTIN, EAN-13, часть серийного номера. Несколько
    условий через пробел или «;» работают как «ИЛИ»."""
    terms = [t for t in re.split(r"[\s;]+", text.strip()) if t]
    if not terms:
        return list(codes)
    return [c for c in codes if any(t in c.full for t in terms)]


def gtin_summary(codes):
    cnt = Counter(c.gtin for c in codes)
    return sorted(cnt.items())


# ---------------------------------------------------------------- отчёт
def make_report(m1, m2, m3, both, only_m2):
    L = []
    for ds in (m1, m2):
        L.append("%s: файлов %d, ячеек/строк %d" % (ds.name, len(ds.files), ds.cells))
        L.append("   кодов найдено %d, повторов удалено %d, брака %d, уникальных %d" %
                 (ds.found, ds.duplicates, ds.bad, len(ds.codes)))
        for f, n, ok, dup, bad in ds.files:
            L.append("   · %s: ячеек %d, кодов %d, повторов %d, брака %d" % (f, n, ok, dup, bad))
    L.append("")
    L.append("Есть и в М1, и в М2: %d" % both)
    L.append("М3 (в М1 есть, в М2 нет): %d" % len(m3))
    L.append("Новые в М2 (нет в М1): %d" % len(only_m2))
    summ = gtin_summary(m3)
    if summ:
        L.append("")
        L.append("М3 по GTIN:")
        for g, n in summ:
            L.append("   %s (EAN-13 %s): %d" % (g, g[1:], n))
    return "\n".join(L)


# ---------------------------------------------------------------- экспорт
def export_txt(codes, path, with_tail):
    with open(path, "w", encoding="utf-8", newline="\r\n") as f:
        for c in codes:
            f.write((c.full if with_tail else c.key) + "\n")


def export_txt_by_gtin(codes, folder):
    """Файл на каждый GTIN: коды без криптохвоста (для загрузки в ЧЗ по группам)."""
    os.makedirs(folder, exist_ok=True)
    groups = OrderedDict()
    for c in sort_codes(codes, "gtin"):
        groups.setdefault(c.gtin, []).append(c)
    for g, lst in groups.items():
        export_txt(lst, os.path.join(folder, "%s.txt" % g), with_tail=False)
    return len(groups)


def export_xlsx(codes, path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Коды"
    ws.append(["GTIN", "EAN-13", "Серийный номер", "Код без криптохвоста", "Код с криптохвостом", "Формат"])
    for c in codes:
        ws.append([c.gtin, c.gtin[1:], c.serial, c.key, c.full, c.kind])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = "@"
            cell.data_type = "s"
    for col, w in zip("ABCDEF", (16, 16, 20, 40, 100, 16)):
        ws.column_dimensions[col].width = w
    ws2 = wb.create_sheet("Для ЧЗ по GTIN")
    ws2.append(["GTIN", "Код без криптохвоста"])
    for c in sort_codes(codes, "gtin"):
        ws2.append([c.gtin, c.key])
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = "@"
            cell.data_type = "s"
    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 44
    ws3 = wb.create_sheet("Сводка")
    ws3.append(["GTIN", "EAN-13", "Количество"])
    for g, n in gtin_summary(codes):
        ws3.append([g, g[1:], n])
    for row in ws3.iter_rows(min_row=2, max_col=2):
        for cell in row:
            cell.number_format = "@"
            cell.data_type = "s"
    wb.save(path)


def export_defects(datasets, path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Брак"
    ws.append(["Набор", "Файл", "Где", "Тип", "Причина", "Фрагмент"])
    for ds in datasets:
        for d in ds.defects:
            ws.append([ds.name, d.source, d.where, "брак" if d.fatal else "поправка", d.reason, d.raw[:200]])
            ws.cell(row=ws.max_row, column=6).data_type = "s"
    for col, w in zip("ABCDEF", (6, 22, 16, 10, 60, 90)):
        ws.column_dimensions[col].width = w
    wb.save(path)
